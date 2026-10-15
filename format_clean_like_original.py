import copy, pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
SRC = '05_FoodAppBusiness_Clean_source.xlsx'  # bản xuất từ bảng Customers của mon.pbix
c = pd.read_excel(SRC, sheet_name='Customers_Clean')
m = pd.read_excel(SRC, sheet_name='Mo_ta_cot').dropna(subset=['Cột'])
m = m[m['Cột'].isin(c.columns)]
log = pd.read_excel(SRC, sheet_name='Cleaning_Log').dropna(subset=['Bước'])
# nhãn khớp báo cáo final
c['RFM_Segment'] = c.RFM_Segment.replace({'New / Promising': 'Promising'})
c['Marital_Status'] = c.Marital_Status.replace({'Married/Partner': 'Married'})
fix = {'Income': 'Thu nhập (nhãn gốc Monthly Income; loại tiền và kỳ tính chưa kiểm chứng được, ghi theo đơn vị tiền tệ của bộ dữ liệu - ĐVTT)',
       'Tenure_Years': 'Tenure_Days ÷ 365 (làm tròn 2 chữ số)',
       'Married': '1 = đã kết hôn (cột Single đã bỏ vì = 1 - Married)',
       'Marital_Status': 'Married hoặc Single (nhãn của Married)',
       'RFM_Segment': 'Champions (R≥4, FM≥4); Loyal Customers (R≥3, FM≥3, còn lại); Promising (R≥4, FM<3); Need Attention (R=3, FM<3); At Risk (R≤2, FM≥3); Hibernating (R≤2, FM<3). Điểm bằng nhau xếp theo Raw_Row',
       'Preferred_Channel': 'Kênh có nhiều đơn nhất; khi bằng nhau ưu tiên Web → Catalog → Store',
       'AOV': 'Total_Spend ÷ Total_Purchases (làm tròn 2 chữ số)', 'Deal_Ratio': 'Deals_Purchases ÷ Total_Purchases (làm tròn 4 chữ số)'}
desc = {r['Cột']: fix.get(r['Cột'], r['Mô tả / cách tính']) for _, r in m.iterrows()}
kind = dict(zip(m['Cột'], m['Loại']))

wb = load_workbook('/home/user/Food-App-Business/05_FoodAppBusiness.xlsx')
ws, wd = wb['Data'], wb['05_FoodAppBusiness']
hs = {k: copy.copy(getattr(ws['A1'], k)) for k in ('font', 'alignment', 'border', 'fill')}
dsty = {k: copy.copy(getattr(ws['A2'], k)) for k in ('font', 'alignment', 'border', 'fill')}
ws.delete_rows(1, ws.max_row)
for k in list(ws.column_dimensions): del ws.column_dimensions[k]
ws.row_dimensions[1].height = 30
cols = list(c.columns)
for j, name in enumerate(cols, 1):
    cell = ws.cell(1, j, name)
    for k, v in hs.items(): setattr(cell, k, copy.copy(v))
for i, row in enumerate(c.astype(object).where(c.notna(), None).values.tolist(), 2):
    for j, v in enumerate(row, 1):
        if isinstance(v, float) and v.is_integer(): v = int(v)
        cell = ws.cell(i, j, v)
        for k, s in dsty.items(): setattr(cell, k, copy.copy(s))
for j, name in enumerate(cols, 1):
    L = max(len(name) * 1.05 + 1.5, max(len(str(x)) for x in c[name].astype(str).head(300)) + 2)
    ws.column_dimensions[get_column_letter(j)].width = round(min(max(7, L), 30), 2)

# ---- sheet mô tả: cùng kiểu với sheet gốc (TNR 14 đoạn giới thiệu, TNR 12 đỏ cho từng biến)
st_intro = {k: copy.copy(getattr(wd['A1'], k)) for k in ('font', 'alignment')}
st_line = {k: copy.copy(getattr(wd['A2'], k)) for k in ('font', 'alignment', 'border', 'fill')}
wd.delete_rows(1, wd.max_row)
intro = ('Bộ dữ liệu Food App sau làm sạch (bảng Customers trong file Power BI mon.pbix), dùng cho báo cáo cuối kỳ. '
         'Mỗi dòng là một khách hàng: 1.992 khách × 66 cột = 25 cột gốc (đã bỏ Single) + 2 cột định danh (CustomerID, Raw_Row) + 39 cột dẫn xuất. '
         'Dữ liệu không ghi loại tiền tệ nên giá trị tiền ghi theo đơn vị tiền tệ của bộ dữ liệu (ĐVTT); chi tiêu là số cộng dồn, không rõ kỳ tính.')
lines = [('i', intro)]
for col in cols:
    tag = '' if kind.get(col, '').startswith('Gốc') else f' [{kind.get(col, "Dẫn xuất")}]'
    lines.append(('l', f'{col} - {desc.get(col, "")}{tag}'))
lines.append(('i', 'Nhật ký làm sạch (Cleaning Log)'))
for r in log.itertuples(index=False):
    step, act, n, left, why = r[1], r[2], r[3], r[4], r[5]
    why = str(why).replace('→ xác suất trùng tự nhiên gần bằng 0; giữ bản ghi đầu tiên', '→ giả định là bản ghi lặp (không có mã khách gốc); giữ bản ghi đầu tiên')
    why = why.replace('33 biến mới', '39 cột dẫn xuất')
    lines.append(('l', f'{step} - {act}: {why} (ảnh hưởng {int(n):,} dòng, còn {int(left):,} dòng)'.replace(',', '.')))
for i, (k, t) in enumerate(lines, 1):
    cell = wd.cell(i, 1, t)
    for a, v in (st_intro if k == 'i' else st_line).items(): setattr(cell, a, copy.copy(v))
    import math; wd.row_dimensions[i].height = 51 if i == 1 else 16 * max(1, math.ceil(len(t) / 105))
wd.title = '05_FoodAppBusiness_Clean'
out = '/home/user/Food-App-Business/05_FoodAppBusiness_Clean.xlsx'
wb.save(out); print('saved', ws.max_row, ws.max_column, wd.max_row, c.RFM_Segment.value_counts().to_dict())
